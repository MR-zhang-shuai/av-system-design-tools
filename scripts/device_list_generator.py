#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
设备清单生成器
==============
功能: 根据设备配置文件生成标准格式的设备清单 Excel
输入: JSON 配置文件
输出: .xlsx 文件（列结构：序号 | 设备名称 | 品牌 | 型号 | 数量 | 单位 | 单价 | 总价 | 参数 | 备注）

使用示例:
    python device_list_generator.py --config sample_config.json --output 设备清单.xlsx

依赖: openpyxl (pip install openpyxl)
"""
import json
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    print("\n需要 openpyxl 库，请运行: pip install openpyxl")
    sys.exit(1)


def create_device_list(config_path, output_path):
    """根据配置文件生成设备清单 Excel"""
    # 读取 JSON 配置
    try:
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
    except FileNotFoundError:
        print("错误: 找不到配置文件 " + config_path)
        return
    except json.JSONDecodeError as e:
        print("错误: 配置文件 JSON 格式错误: " + str(e))
        return

    devices = config.get('devices', [])
    if not devices:
        print("警告: 配置文件中没有设备数据，将生成空清单")

    # 创建工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "设备清单"

    # 表头定义
    headers = ["序号", "设备名称", "品牌", "型号", "数量", "单位", "单价", "总价", "参数", "备注"]
    header_font = Font(name="微软雅黑", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 写入设备数据
    data_font = Font(name="微软雅黑", size=9)
    data_align = Alignment(vertical="center", wrap_text=True)

    for row_idx, device in enumerate(devices, 2):
        row_data = [
            row_idx - 1,
            device.get('name', ''),
            device.get('brand', ''),
            device.get('model', ''),
            device.get('qty', 1),
            device.get('unit', '台'),
            device.get('unit_price', 0),
            device.get('total_price', 0),
            device.get('params', ''),
            device.get('remark', '')
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border

    # 设置列宽
    col_widths = [6, 25, 15, 25, 8, 6, 12, 12, 40, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    # 保存文件
    try:
        wb.save(output_path)
        print("设备清单已生成: " + output_path)
        print("共 " + str(len(devices)) + " 项设备")
    except Exception as e:
        print("错误: 保存文件失败: " + str(e))


def main():
    import argparse
    parser = argparse.ArgumentParser(description="设备清单生成器")
    parser.add_argument('--config', required=True, help='配置文件路径 (JSON)')
    parser.add_argument('--output', default='设备清单.xlsx', help='输出 Excel 文件路径')
    args = parser.parse_args()
    create_device_list(args.config, output_path=args.output)


if __name__ == '__main__':
    main()